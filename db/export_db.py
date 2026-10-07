# -*- coding: utf-8 -*-
"""tarot_data.json → 검수용 엑셀(쏠타로_문장DB.xlsx)

실행: python3 export_db.py ../02_데이터/tarot_data.json 쏠타로_문장DB.xlsx
"""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from db_schema import SHEETS, STATUSES

SRC, OUT = sys.argv[1], sys.argv[2]
# 선택: 세 번째 인자로 점검 결과 JSON(review_log.json)을 주면 '점검이력' 시트와 검수 메모를 채움
REVIEW = json.load(open(sys.argv[3], encoding="utf-8")) if len(sys.argv) > 3 else None
MEMO = {}
if REVIEW:
    for e in REVIEW["edits"]:
        MEMO[e["id"]] = f"[Claude 점검 · 고침] {e['prob']} (전: {e['old']})"
    for e in REVIEW["suggest"]:
        MEMO[e["id"]] = f"[Claude 점검 · 제안] 직접 정해 주신 문구라 그대로 둠. 제안: {e['new']} — {e['prob']}"
    for e in REVIEW["hold"]:
        MEMO[e["id"]] = f"[Claude 점검 · {e['decision']}] {e['prob']}"
    for i, m in REVIEW.get("notes", {}).items():
        MEMO[i] = (MEMO[i] + " / " if i in MEMO else "") + m
d = json.load(open(SRC, encoding="utf-8"))

FONT = "Malgun Gothic"  # 한글 문장 위주라 맑은 고딕(윈도우 기본 한글 글꼴)
NAVY, GOLD, CREAM, INPUT = "012F53", "E3C68E", "FEF3DE", "FFF4B8"
f_body = Font(name=FONT, size=10)
f_id = Font(name=FONT, size=8, color="8A94A0")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_h2 = Font(name=FONT, size=12, bold=True, color=NAVY)
fill_head = PatternFill("solid", fgColor=NAVY)
fill_input = PatternFill("solid", fgColor=INPUT)
fill_cream = PatternFill("solid", fgColor=CREAM)
thin = Side(style="thin", color="D9D2C3")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="top", wrap_text=True)

HEAD = ["ID (수정 금지)", "카드 · 분류", "항목", "현재 문장 (초안)", "상태", "수정안", "검수 메모", "최종 글자 수", "분량"]
WIDTH = [24, 22, 26, 62, 10, 56, 30, 9, 10]

wb = Workbook()
guide = wb.active
guide.title = "읽어주세요"

counts = {}
for name, what, where, (lo, hi), fn, stage in SHEETS:
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = GOLD if name == "깊이읽기" else NAVY
    for i, (h, w) in enumerate(zip(HEAD, WIDTH), start=1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, box
        ws.column_dimensions[c.column_letter].width = w
    ws.row_dimensions[1].height = 30
    r = 1
    for rid, label, item, text in fn(d):
        r += 1
        vals = [rid, label, item, text, "초안", None, MEMO.get(rid),
                f'=LEN(IF(F{r}<>"",F{r},D{r}))',
                f'=IF(OR(H{r}<{lo},H{r}>{hi}),"분량 확인","")']
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=r, column=i, value=v)
            c.font, c.border = (f_id if i == 1 else f_body), box
            c.alignment = center if i in (5, 8, 9) else wrap
            if i in (5, 6, 7):
                c.fill = fill_input
        lines = max(-(-len(text) // 29), -(-len(str(label)) // 10), -(-len(str(item)) // 12), 1)
        ws.row_dimensions[r].height = 14 * lines + 6  # 엑셀이 줄높이를 자동으로 안 맞출 때를 대비
    counts[name] = r - 1
    dv = DataValidation(type="list", formula1='"' + ",".join(STATUSES) + '"', allow_blank=False,
                        error="초안 · 검수중 · 수정요청 · 확정 · 삭제 중에서 골라 주세요.", errorTitle="상태")
    ws.add_data_validation(dv)
    dv.add(f"E2:E{r}")
    rng = f"E2:E{r}"
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"확정"'], fill=PatternFill("solid", fgColor="CFE8C9")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"수정요청"'], fill=PatternFill("solid", fgColor="F8D2B0")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"검수중"'], fill=PatternFill("solid", fgColor="D6E4F2")))
    ws.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=['"삭제"'], fill=PatternFill("solid", fgColor="D0D0D0")))
    ws.conditional_formatting.add(f"I2:I{r}", FormulaRule(formula=[f'I2="분량 확인"'], font=Font(name=FONT, color="B4441C", bold=True)))
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = f"A1:I{r}"
    ws.sheet_view.zoomScale = 110

# ---- 읽어주세요 ----
g = guide
g.sheet_properties.tabColor = GOLD
for col, w in zip("ABCDEFGHIJ", [16, 30, 34, 18, 10, 10, 10, 10, 11, 11]):
    g.column_dimensions[col].width = w
row = 1
def put(r, c, v, font=f_body, fill=None, align=wrap, border=None):
    cell = g.cell(row=r, column=c, value=v)
    cell.font, cell.alignment = font, align
    if fill: cell.fill = fill
    if border: cell.border = border
    return cell

put(1, 1, "쏠 타로 하우스 — 미리 써 두는 문장 DB", f_title, align=Alignment(vertical="center"))
g.row_dimensions[1].height = 30
put(2, 1, f"기준 데이터: tarot_data.json v{d['version']} ({d['date']}). 모든 문장은 Claude 초안이며 상담사 검수 전입니다. 결과 리포트는 이 문장들을 서버가 조립하고, AI는 고민 맞춤 한 줄·질문·이야기만 짧게 씁니다.")
g.merge_cells("A2:J2"); g.row_dimensions[2].height = 32

put(4, 1, "쓰는 법", f_h2, align=Alignment(vertical="bottom"))
steps = [
    "1. 시트 하나를 열고 '현재 문장'을 읽습니다. 노란 칸(상태 · 수정안 · 검수 메모)만 고쳐 주세요.",
    "2. 그대로 써도 되면 상태를 '확정'으로 바꿉니다. 고칠 문장은 '수정안'에 새 문장을 쓰고 상태를 '확정'으로 둡니다.",
    "3. 직접 고치기 어렵거나 의견만 남길 때는 상태 '수정요청' + '검수 메모'에 이유를 적어 주세요.",
    "4. '현재 문장'과 ID 칸은 고치지 마세요. 수정안이 있으면 수정안이, 없으면 현재 문장이 최종 문장이 됩니다.",
    "5. {name} · {nickname} 같은 중괄호는 닉네임 등이 들어갈 자리라 그대로 남겨 주세요.",
    "6. '분량 확인'이 뜨면 기준 글자 수를 벗어난 문장이에요. 꼭 맞출 필요는 없지만 화면에서 너무 길거나 짧을 수 있어요.",
    "7. 다 끝나면 파일을 돌려주세요. 개발 쪽에서 import_db.py로 tarot_data.json에 한 번에 반영합니다(확정만 반영하는 옵션 있음).",
]
for i, s in enumerate(steps):
    put(5 + i, 1, s); g.merge_cells(start_row=5 + i, start_column=1, end_row=5 + i, end_column=10)
    g.row_dimensions[5 + i].height = 30

r0 = 13
put(r0, 1, "색 · 상태", f_h2, align=Alignment(vertical="bottom"))
legend = [("노란 칸", "검수자가 쓰는 칸 (상태 · 수정안 · 검수 메모)", fill_input),
          ("확정", "이대로 결과에 쓴다 (수정안이 있으면 수정안으로)", PatternFill("solid", fgColor="CFE8C9")),
          ("수정요청", "고쳐야 하지만 수정안을 아직 못 씀 — 메모 확인", PatternFill("solid", fgColor="F8D2B0")),
          ("검수중", "읽는 중, 나중에 다시 봄", PatternFill("solid", fgColor="D6E4F2")),
          ("초안", "아직 안 봄 (처음 상태)", None),
          ("삭제", "이 문장을 빼고 싶음 — 개발 쪽과 구조를 함께 바꿔야 해서 자동 반영되지 않음", PatternFill("solid", fgColor="D0D0D0"))]
for i, (k, v, f) in enumerate(legend):
    put(r0 + 1 + i, 1, k, Font(name=FONT, size=10, bold=True), f, center, box)
    put(r0 + 1 + i, 2, v); g.merge_cells(start_row=r0 + 1 + i, start_column=2, end_row=r0 + 1 + i, end_column=10)

r1 = r0 + 8
put(r1, 1, "작성 예시 (실제 시트와 같은 칸 구성)", f_h2, align=Alignment(vertical="bottom"))
ex_head = ["ID", "카드 · 분류", "항목", "현재 문장 (초안)", "상태", "수정안", "검수 메모"]
ex_val = ["cards.12.positions.now", "12 바위 위의 쉼 · 물범", "지금의 마음 · 정방향",
          "몸도 마음도 쉬고 싶다는 신호를 보내고 있어요. 열심히 해 왔는데 생각만큼 앞으로 나아가지 않는 느낌이 들 수 있어요.",
          "확정", "몸도 마음도 잠시 쉬어 가자는 신호를 보내고 있어요. 열심히 달려왔는데 제자리인 것 같은 느낌이 들 수 있어요.",
          "'나아가지 않는'이 실패처럼 들려서 '제자리인 것 같은'으로 부드럽게"]
spans = [(1, 1), (2, 2), (3, 3), (4, 5), (6, 6), (7, 8), (9, 10)]
for (a, b), h, v in zip(spans, ex_head, ex_val):
    put(r1 + 1, a, h, f_head, fill_head, center, box)
    put(r1 + 2, a, v, f_id if a == 1 else f_body, fill_input if h in ("상태", "수정안", "검수 메모") else None, wrap, box)
    if b > a:
        g.merge_cells(start_row=r1 + 1, start_column=a, end_row=r1 + 1, end_column=b)
        g.merge_cells(start_row=r1 + 2, start_column=a, end_row=r1 + 2, end_column=b)
g.row_dimensions[r1 + 2].height = 96

r2 = r1 + 4
put(r2, 1, "진행 현황 (자동 계산)", f_h2, align=Alignment(vertical="bottom"))
ph = ["시트", "무엇인가", "결과에서 쓰이는 곳", "단계", "문장 수", "확정", "수정요청", "초안", "확정 비율", "분량 확인"]
for i, h in enumerate(ph, start=1):
    put(r2 + 1, i, h, f_head, fill_head, center, box)
first = r2 + 2
for j, (name, what, where, _, _, stage) in enumerate(SHEETS):
    rr = first + j
    q = f"'{name}'"
    vals = [name, what, where, stage,
            f"=COUNTA({q}!A:A)-1",
            f'=COUNTIF({q}!E:E,"확정")',
            f'=COUNTIF({q}!E:E,"수정요청")',
            f'=COUNTIF({q}!E:E,"초안")',
            f"=IF(E{rr}=0,0,F{rr}/E{rr})",
            f'=COUNTIF({q}!I:I,"분량 확인")']
    for i, v in enumerate(vals, start=1):
        c = put(rr, i, v, f_body, fill_cream if name == "깊이읽기" else None, center if i >= 5 else wrap, box)
        if i == 9: c.number_format = "0%"
        if i == 1: c.hyperlink = f"#{q}!A1"; c.font = Font(name=FONT, size=10, color="1F4E8C", underline="single")
last = first + len(SHEETS) - 1
tr = last + 1
put(tr, 1, "합계", Font(name=FONT, size=10, bold=True), None, center, box)
for i in range(2, 5):
    put(tr, i, None, f_body, None, wrap, box)
for i, col in zip(range(5, 11), "EFGHIJ"):
    v = f"=IF(E{tr}=0,0,F{tr}/E{tr})" if col == "I" else f"=SUM({col}{first}:{col}{last})"
    c = put(tr, i, v, Font(name=FONT, size=10, bold=True), None, center, box)
    if col == "I": c.number_format = "0%"
put(tr + 1, 1, "깊이읽기 시트(금색 탭)는 이번에 새로 쓴 2단계 초안입니다. 역방향은 위치별 공통 단락 3개를 22장 모두에 씁니다.", Font(name=FONT, size=9, color="5A6470"))
g.merge_cells(start_row=tr + 1, start_column=1, end_row=tr + 1, end_column=10)
g.freeze_panes = "A4"

if REVIEW:
    ws = wb.create_sheet("점검이력", 1)
    ws.sheet_properties.tabColor = "B4441C"
    hd = ["ID", "구분", "처리", "무엇이 문제였나", "전", "후 (현재 문장)"]
    for i, (h, w) in enumerate(zip(hd, [26, 8, 16, 40, 52, 52]), start=1):
        c = ws.cell(row=1, column=i, value=h); c.font, c.fill, c.alignment, c.border = f_head, fill_head, center, box
        ws.column_dimensions[c.column_letter].width = w
    order = {"우려": 0, "오류": 1, "다듬기": 2}
    items = [(e["id"], e["kind"], e["decision"], e["prob"], e["old"], e["new"]) for e in REVIEW["edits"]]
    items += [(e["id"], e["kind"], "제안만(문구 유지)", e["prob"], "", e["new"]) for e in REVIEW["suggest"]]
    items += [(e["id"], e["kind"], e["decision"], e["prob"], "", "") for e in REVIEW["hold"]]
    items.sort(key=lambda x: (order.get(x[1], 3), x[0]))
    for r, it in enumerate(items, start=2):
        for i, v in enumerate(it, start=1):
            c = ws.cell(row=r, column=i, value=v); c.font = f_id if i == 1 else f_body; c.alignment = wrap; c.border = box
        ws.row_dimensions[r].height = 14 * max(1, -(-max(len(it[3]), len(it[4]), len(it[5])) // 24)) + 6
    ws.freeze_panes = "B2"; ws.auto_filter.ref = f"A1:F{len(items) + 1}"

wb.save(OUT)
print("저장:", OUT, counts, "합계", sum(counts.values()))
